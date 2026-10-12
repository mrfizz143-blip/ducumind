from __future__ import annotations

from pathlib import Path

from .base import BaseProcessor, Block, ExtractionResult


class DOCXProcessor(BaseProcessor):
    name = "docx"

    @staticmethod
    def supports(filename: str) -> bool:
        return filename.lower().endswith(".docx")

    def process(self, file_path: Path, original_name: str, media_dir: Path) -> ExtractionResult:
        from docx import Document

        doc = Document(str(file_path))
        result = ExtractionResult(title=Path(original_name).stem)

        core = doc.core_properties
        if core.title:
            result.title = core.title
        result.author = core.author or ""
        result.metadata = {
            "author": core.author or "",
            "subject": core.subject or "",
            "created": str(core.created) if core.created else "",
            "modified": str(core.modified) if core.modified else "",
            "keywords": core.keywords or "",
        }

        paragraphs = []
        current_section = ""
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name or "").lower() if p.style else ""
            is_heading = style.startswith("heading") or style.startswith("title")
            if is_heading:
                current_section = text
            block_type = "heading" if is_heading else "paragraph"
            result.blocks.append(Block(content=text, section=current_section, block_type=block_type))
            paragraphs.append(text)

        for tbl in doc.tables:
            rows = []
            for row in tbl.rows:
                cells = [c.text.strip() for c in row.cells]
                rows.append(" | ".join(cells))
            if rows:
                result.blocks.append(
                    Block(content="\n".join(rows), section=current_section, block_type="table")
                )
                paragraphs.append("Table:\n" + "\n".join(rows))

        result.text = "\n\n".join(paragraphs)
        result.toc = [
            {"level": 1, "title": b.content, "page": 0}
            for b in result.blocks
            if b.block_type == "heading"
        ]
        return result